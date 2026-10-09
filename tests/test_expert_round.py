"""The expert-review round: one blinded workbook for three reviews, and its import back into each review's format."""
import importlib.util
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bioevidence_validator import audit, review

ROOT = Path(__file__).resolve().parents[1]
ROUND = ROOT / "evaluation" / "expert_round"
KIT = ROOT / "evaluation" / "clinvar_review"
sys.path.insert(0, str(ROUND))
sys.path.insert(0, str(KIT))


def load(folder, name):
    spec = importlib.util.spec_from_file_location(f"er_{name}", folder / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


mk, imp, prepare = load(ROUND, "make_round"), load(ROUND, "import_round"), load(KIT, "prepare_packets")


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("round")
    prepare.prepare(tmp / "kit", salt="test-salt", calibration=20)
    assert mk.main(["--clinvar-packet", str(tmp / "kit" / "reviewer_packet"), "--clinvar-key",
                    str(tmp / "kit" / "maintainer" / "key.json"), "--seed", "7", "--output", str(tmp / "round")]) == 0
    return tmp


def fill(src, dst, answers, who=("R1", "curator, 5 years", "2026-10-15"), rows=None):
    book = load_workbook(src)
    start = book["Start here"]
    start["B4"], start["B5"], start["B6"] = who
    for name, values in answers.items():
        sheet = book[mk.PARTS[name]["sheet"]]
        columns = [i for i, (_, _, c) in enumerate(mk.COLUMNS[name], start=1) if isinstance(c, dict)]
        last = sheet.max_row if rows is None else mk.HEADER_ROW + rows
        for row in range(mk.HEADER_ROW + 1, last + 1):
            for column, value in zip(columns, values, strict=False):
                sheet.cell(row=row, column=column, value=value)
    book.save(dst)
    return dst


def test_workbook_is_blinded_and_every_answer_is_a_dropdown(built):
    book = load_workbook(built / "round" / mk.WORKBOOK)
    assert book.sheetnames == ["Start here", "Lists", *(p["sheet"] for p in mk.PARTS.values())]
    assert book["Lists"].sheet_state == "hidden"
    key = json.loads((built / "round" / "maintainer" / "round_key.json").read_text(encoding="utf-8"))
    sizes = {name: len(items) for name, items in key["items"].items()}
    assert sizes == {"A": mk.CLINVAR_CASES, "B": 60, "C": audit.sample_size(0.05) + 10}
    hidden = ("claude", "gpt", "gemini", "validator", "divergent", "review_required", "bioevidence", "llm-feedback")
    for name, part in mk.PARTS.items():
        sheet = book[part["sheet"]]
        text = " ".join(str(c.value) for row in sheet.iter_rows(max_col=4) for c in row if c.value).lower()
        if name != "B":  # paper quotes may say anything
            assert not any(word in text for word in hidden)
        questions = [t for t, _, c in mk.COLUMNS[name] if isinstance(c, dict)]
        assert len(sheet.data_validations.dataValidation) == len(questions)
        assert all(dv.formula1.startswith("=Lists!") for dv in sheet.data_validations.dataValidation)
    clinvar = json.loads((built / "kit" / "maintainer" / "key.json").read_text(encoding="utf-8"))
    roles = [clinvar["cases"][code]["role"] for code in key["items"]["A"]]
    assert roles.count("divergent") == roles.count("control") == mk.CLINVAR_CASES // 2


def test_returned_workbook_imports_into_each_review(built, tmp_path):
    filled = fill(built / "round" / mk.WORKBOOK, tmp_path / "R1.xlsx",
                  {"A": ["yes", "usable as is", "needs a curator's check", "not usable"],
                   "B": ["no", "Opposite direction"], "C": ["unsure", "needs a curator's check"]})
    out = tmp_path / "out"
    args = [str(filled), "--key", str(built / "round" / "maintainer" / "round_key.json"),
            "--clinvar-key", str(built / "kit" / "maintainer" / "key.json"), "--output", str(out)]
    assert imp.main(args) == 0
    clinvar = review.load_annotations(out / "clinvar_annotations.csv")
    assert len(clinvar) == 3 * mk.CLINVAR_CASES and {a["reviewer_id"] for a in clinvar} == {"R1"}
    assert {a["requested_use"]: a["admission_label"] for a in clinvar} == {
        "research_summary": "admitted", "clinical_reference": "review_required", "expert_reference": "rejected"}
    extraction = (out / "extraction_labels.csv").read_text(encoding="utf-8").splitlines()
    assert len(extraction) == 61 and extraction[1].split(",")[2:4] == ["no", "EXT-2"]
    singlecell = review.load_annotations(out / "singlecell_annotations.csv")
    manifest = json.loads((built / "round" / "maintainer" / "singlecell_audit_manifest.json").read_text(encoding="utf-8"))
    report = audit.audit_score(manifest, review.resolve(singlecell, [], min_reviewers=1), singlecell)
    assert report["routes"]["admitted"]["audited"] == audit.sample_size(0.05)
    with pytest.raises(ValueError, match="already holds"):
        imp.main(args)  # the same reviewer twice


def test_partial_and_invalid_answers(built, tmp_path):
    key = ["--key", str(built / "round" / "maintainer" / "round_key.json"), "--output", str(tmp_path / "out")]
    partial = fill(built / "round" / mk.WORKBOOK, tmp_path / "p.xlsx", {"C": ["yes", "usable as is"]}, rows=5)
    assert imp.main([str(partial), *key]) == 0  # only part C, five rows; no ClinVar key needed
    assert len(review.load_annotations(tmp_path / "out" / "singlecell_annotations.csv")) == 5
    half = fill(built / "round" / mk.WORKBOOK, tmp_path / "h.xlsx", {"A": ["yes"]}, rows=1)
    with pytest.raises(ValueError, match="unanswered"):
        imp.main([str(half), *key, "--reviewer-id", "R2"])
    wrong = fill(built / "round" / mk.WORKBOOK, tmp_path / "w.xlsx", {"C": ["maybe", "usable as is"]}, rows=1)
    with pytest.raises(ValueError, match="not an option"):
        imp.main([str(wrong), *key, "--reviewer-id", "R3"])
    anonymous = fill(built / "round" / mk.WORKBOOK, tmp_path / "a.xlsx", {"C": ["yes", "usable as is"]},
                     who=(None, None, None), rows=1)
    with pytest.raises(ValueError, match="ID and one-line qualification"):
        imp.main([str(anonymous), *key])
