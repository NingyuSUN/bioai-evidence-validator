"""ClinVar expert-review kit: selection, blinding, workbook round trip and scoring."""
import csv
import importlib.util
import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bioevidence_validator import review
from bioevidence_validator.cli import main

ROOT = Path(__file__).resolve().parents[1]
KIT = ROOT / "evaluation" / "clinvar_review"


def load(name):
    spec = importlib.util.spec_from_file_location(f"kit_{name}", KIT / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


prepare, importer = load("prepare_packets"), load("import_sheets")


@pytest.fixture(scope="module")
def kit(tmp_path_factory):
    out = tmp_path_factory.mktemp("kit") / "packet"
    counts = prepare.prepare(out, salt="fixed-test-salt", calibration=20)
    return out, counts, json.loads((out / "maintainer" / "key.json").read_text(encoding="utf-8"))


def test_selection_is_every_divergence_plus_matched_controls(kit):
    _, counts, key = kit
    with (ROOT / "examples/clinvar_germline/results/divergences.csv").open(encoding="utf-8") as handle:
        divergent = {r["variation_id"] for r in csv.DictReader(handle)}
    selected = {c["variation_id"] for c in key["cases"].values() if c["role"] == "divergent"}
    assert selected == divergent and counts["divergent"] == len(divergent) == 95
    by_role = {role: sorted(c["stratum"] for c in key["cases"].values() if c["role"] == role)
               for role in ("divergent", "control")}
    assert by_role["divergent"] == by_role["control"]  # controls matched stratum for stratum
    assert counts["calibration"] == 20 and sum(c["split"] == "development" for c in key["cases"].values()) == 20


def test_reviewer_packet_is_blinded(kit):
    out, _, key = kit
    packet = out / "reviewer_packet"
    assert sorted(p.name for p in packet.iterdir()) == ["RUBRIC.md", "evidence.csv", "labels.csv", "review_workbook.xlsx"]
    cells = []
    for name in ("labels.csv", "evidence.csv"):
        with (packet / name).open(encoding="utf-8") as handle:
            cells += [v for r in csv.reader(handle) for v in r]
    book = load_workbook(packet / "review_workbook.xlsx", read_only=True)
    cells += [str(v) for sheet in book for r in sheet.iter_rows(values_only=True) for v in r if v is not None]
    text = "\n".join(cells)
    ids = {c["variation_id"] for c in key["cases"].values()}
    assert not ids & set(cells), "a VariationID appears as a cell value"
    assert not re.search(r"SCV\d+|multiple submitters|no conflicts|conflicting interpretations|BEV0\d\d", text)
    assert "fixed-test-salt" not in text and not re.search(r"\badmitted\b.*\bvalidator\b", text)
    for status in ("criteria provided, single submitter", "no assertion criteria provided"):
        assert status not in text  # per-submission wording is translated, never ClinVar's aggregate phrases


def fill(path, key, flip_every=0):
    book = load_workbook(path)
    sheet = book["Labels"]
    header = [c.value for c in sheet[1]]
    col = {name: header.index(name) + 1 for name in header}
    for n, cells in enumerate(sheet.iter_rows(min_row=2), start=1):
        case = key["cases"][cells[0].value]
        sheet.cell(cells[0].row, col["statement_label"], "correct")
        for use in importer.USES:
            label = case["validator"][use]
            if flip_every and n % flip_every == 0 and use == "research_summary":
                label = "admitted" if label != "admitted" else "review_required"
            sheet.cell(cells[0].row, col[use], label)
            sheet.cell(cells[0].row, col[f"{use}_rationale"], "synthetic reviewer")
        sheet.cell(cells[0].row, col["sources_consulted"], "ACMG/AMP 2015; gene background")
    book.save(path)
    return path


def test_workbook_round_trip_adjudication_and_scoring(kit, tmp_path):
    out, _, key = kit
    workbook = out / "reviewer_packet" / "review_workbook.xlsx"
    annotations = tmp_path / "annotations.csv"
    _import(_copy_fill(workbook, tmp_path / "R1.xlsx", key, 0), "R1", key, annotations)
    _import(_copy_fill(workbook, tmp_path / "R2.xlsx", key, 9), "R2", key, annotations)  # R2 differs on every 9th case
    rows = review.load_annotations(annotations)
    assert len(rows) == 2 * 190 * 3
    before = annotations.read_bytes()
    with pytest.raises(ValueError, match="twice"):
        _import(tmp_path / "R1.xlsx", "R1", key, annotations)
    assert annotations.read_bytes() == before and not list(tmp_path.glob("*.tmp.csv"))  # a failed import changes nothing

    sheet, incomplete = review.adjudication_sheet(rows)
    assert incomplete == [] and len(sheet) == 21  # every 9th of 190 cases, research_summary only
    for r in sheet:
        r.update(mapping_label="correct", admission_label=key["cases"][_code(key, r["case_id"])]["validator"][r["requested_use"]],
                 adjudicator_id="A1", adjudicated_at="2026-10-20", mapping_rationale="-", admission_rationale="-",
                 evidence_refs_json='["packet"]')
    review.write_csv(tmp_path / "adj.csv", sheet, review.ADJUDICATION_COLUMNS)
    final = review.resolve(rows, review.load_adjudications(tmp_path / "adj.csv", rows))
    report = review.score(final, review.load_predictions(out / "maintainer" / "predictions.csv"))
    validator = report["methods"]["validator"]["all"]
    assert validator["n"] == 170 * 3 and validator["exact_matches"] == validator["n"]
    assert set(report["methods"]["ncbi_review_status"]) == {"all", "divergent", "control"}
    assert main(["review", "check", str(annotations), "--adjudications", str(tmp_path / "adj.csv")]) == 0


def _code(key, case_id):
    return next(code for code, c in key["cases"].items() if c["case_id"] == case_id)


def _copy_fill(source, target, key, flip_every):
    target.write_bytes(source.read_bytes())
    return fill(target, key, flip_every)


def _import(path, reviewer, key, output):
    rows, _ = importer.convert(importer.read_labels(path), key, reviewer_id=reviewer,
                               qualification="synthetic test reviewer", annotated_at="2026-10-15")
    return importer.append(output, rows)


def test_partial_returns_import_and_half_filled_rows_fail(kit, tmp_path):
    out, _, key = kit
    labels = importer.read_labels(out / "reviewer_packet" / "labels.csv")
    first = {**labels[0], "statement_label": "Correct", "research_summary": "Review required",
             "clinical_reference": "rejected", "expert_reference": "rejected"}
    rows, skipped = importer.convert([first, *labels[1:]], key, reviewer_id="R9", qualification="q",
                                     annotated_at="2026-10-15")
    assert len(rows) == 3 and skipped == 189 and rows[0]["admission_label"] == "review_required"
    with pytest.raises(ValueError, match="missing or invalid"):
        importer.convert([{**labels[0], "statement_label": "correct"}], key, reviewer_id="R9", qualification="q",
                         annotated_at="2026-10-15")
