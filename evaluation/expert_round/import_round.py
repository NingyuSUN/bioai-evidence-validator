"""Read a returned expert-review workbook back into each review's own format (the counterpart of make_round.py).

    uv run --frozen python evaluation/expert_round/import_round.py returned_R1.xlsx \
        --key artifacts/expert-round-1/maintainer/round_key.json \
        --clinvar-key artifacts/clinvar-review/maintainer/key.json \
        --reviewer-id R1 --output artifacts/expert-round-1/returned

Each part the expert filled in is appended to its own file; blank rows (parts outside their field) are skipped:
- Part A -> `clinvar_annotations.csv`, the protocol annotation format of the ClinVar review kit
  (`bioevidence review agreement`, `adjudication-sheet`, `score`);
- Part B -> `extraction_labels.csv`, the labelling format of the extraction expert sample;
- Part C -> `singlecell_annotations.csv`, protocol annotations for `bioevidence review audit-score` with
  `maintainer/singlecell_audit_manifest.json`.

A half-answered row, or an answer that is not one of the dropdown's options, is an error; nothing is written then.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(REPO / "evaluation" / "clinvar_review"))
import import_sheets  # noqa: E402
import make_round as mk  # noqa: E402

from bioevidence_validator import review  # noqa: E402

EXTRACTION_COLUMNS = ["sample_id", "pmid", "claim_correct", "error_codes", "rationale", "reviewer_id",
                      "reviewer_qualification", "annotated_at", "minutes_spent"]


def read_sheet(sheet) -> list[dict[str, Any]]:
    header = [c.value for c in sheet[mk.HEADER_ROW]]
    rows = []
    for values in sheet.iter_rows(min_row=mk.HEADER_ROW + 1, values_only=True):
        row = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, values, strict=False) if h}
        if row.get("Item"):
            rows.append(row)
    return rows


def answers(row: dict[str, str], name: str) -> dict[str, str]:
    """The answer columns of one row (questions only), checked against the dropdown options."""
    out = {}
    for title, _, choices in mk.COLUMNS[name]:
        if isinstance(choices, dict):
            value = row.get(title, "")
            if value and value not in choices:
                raise ValueError(f"{row['Item']}: {value!r} is not an option for {title!r}")
            out[title] = value
    return out


def reviewer(book, override: dict[str, str]) -> dict[str, str]:
    start = book["Start here"]
    found = {"id": start["B4"].value, "qualification": start["B5"].value, "date": start["B6"].value}
    found = {k: ("" if v is None else str(v).strip()) for k, v in found.items()}
    found.update({k: v for k, v in override.items() if v})
    if not found["id"] or not found["qualification"]:
        raise ValueError("The reviewer's ID and one-line qualification are needed (on the Start sheet or as options)")
    date = found["date"][:10] or dt.date.today().isoformat()
    try:
        dt.date.fromisoformat(date)
    except ValueError:
        raise ValueError(f"Date finished must be YYYY-MM-DD, not {found['date']!r}") from None
    found["date"] = date
    return found


def convert(book, key: dict, clinvar_key: dict | None, who: dict[str, str]) -> dict[str, list[dict[str, str]]]:
    out: dict[str, list[dict[str, str]]] = {"A": [], "B": [], "C": []}
    errors: list[str] = []
    for name, part in mk.PARTS.items():
        for row in read_sheet(book[part["sheet"]]):
            try:
                given = answers(row, name)
            except ValueError as exc:
                errors.append(str(exc))
                continue
            if not any(given.values()):
                continue  # not reviewed: a part outside this expert's field
            questions = [t for t, _, c in mk.COLUMNS[name] if isinstance(c, dict)]
            needed = questions if name != "B" else questions[:1]  # B's Q2 only applies when Q1 is "no"
            missing = [q for q in needed if not given[q]]
            if missing:
                errors.append(f"{row['Item']}: unanswered {missing}")
                continue
            source = key["items"][name].get(row["Item"])
            if source is None:
                errors.append(f"{row['Item']}: not an item of this round")
                continue
            comment = row.get("Comment (optional)", "")
            if name == "A":
                q = [given[t] for t in questions]
                out["A"].append({"case_code": source["case_code"], "statement_label": mk.STATEMENT[q[0]],
                                 **{use: mk.USE[v] for use, v in zip(import_sheets.USES, q[1:], strict=True)},
                                 "statement_rationale": comment, "sources_consulted": "", "later_information_seen": "no"})
            elif name == "B":
                verdict, problem = given[questions[0]], given[questions[1]]
                out["B"].append({"sample_id": source["sample_id"], "pmid": source["pmid"], "claim_correct": verdict,
                                 "error_codes": mk.PROBLEMS[problem] if problem and verdict == "no" else "",
                                 "rationale": comment, "reviewer_id": who["id"],
                                 "reviewer_qualification": who["qualification"], "annotated_at": who["date"],
                                 "minutes_spent": ""})
            else:
                verdict, use = given[questions[0]], given[questions[1]]
                out["C"].append({**{c: "" for c in review.ANNOTATION_COLUMNS}, **source,
                                 "annotation_id": f"{who['id']}:{row['Item']}", "group_id": source["case_id"],
                                 "reviewer_id": who["id"], "reviewer_qualification": who["qualification"],
                                 "annotated_at": who["date"], "mapping_label": mk.STATEMENT[verdict],
                                 "mapping_rationale": comment, "admission_label": mk.USE[use],
                                 "admission_rationale": comment, "evidence_refs_json": json.dumps(["expert review round 1"])})
    if errors:
        raise ValueError("Cannot import:\n" + "\n".join(errors[:30]))
    if out["A"]:
        if clinvar_key is None:
            raise ValueError("Part A was answered: --clinvar-key is needed to import it")
        out["A"], _ = import_sheets.convert(out["A"], clinvar_key, reviewer_id=who["id"],
                                            qualification=who["qualification"], annotated_at=who["date"])
    return out


def append_csv(path: Path, rows: list[dict[str, str]], columns: list[str], key_columns: tuple[str, ...]) -> None:
    existing = []
    if path.exists():
        with path.open(encoding="utf-8", newline="") as handle:
            existing = list(csv.DictReader(handle))
    seen = {tuple(r[c] for c in key_columns) for r in existing}
    clash = [r for r in rows if tuple(r[c] for c in key_columns) in seen]
    if clash:
        raise ValueError(f"{path.name} already holds this reviewer's answer for {clash[0][key_columns[0]]}")
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, lineterminator="\n")
        writer.writeheader()
        writer.writerows(existing + rows)


def main(argv: list[str] | None = None) -> int:
    from openpyxl import load_workbook

    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--key", type=Path, required=True, help="maintainer/round_key.json from make_round.py")
    parser.add_argument("--clinvar-key", type=Path, help="The ClinVar kit's maintainer key (for part A)")
    parser.add_argument("--reviewer-id", default="", help="Overrides the Start sheet")
    parser.add_argument("--qualification", default="", help="Overrides the Start sheet")
    parser.add_argument("--annotated-at", default="", help="YYYY-MM-DD; overrides the Start sheet")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    book = load_workbook(args.workbook, read_only=False)
    who = reviewer(book, {"id": args.reviewer_id, "qualification": args.qualification, "date": args.annotated_at})
    key = json.loads(args.key.read_text(encoding="utf-8"))
    clinvar_key = json.loads(args.clinvar_key.read_text(encoding="utf-8")) if args.clinvar_key else None
    parts = convert(book, key, clinvar_key, who)
    args.output.mkdir(parents=True, exist_ok=True)
    # Refuse a second import of the same reviewer before anything is written, so no part is imported alone.
    for name, file, columns in (("A", "clinvar_annotations.csv", ("annotation_id",)),
                                ("B", "extraction_labels.csv", ("sample_id", "reviewer_id")),
                                ("C", "singlecell_annotations.csv", ("case_id", "reviewer_id"))):
        if parts[name] and (args.output / file).exists():
            with (args.output / file).open(encoding="utf-8", newline="") as handle:
                seen = {tuple(r[c] for c in columns) for r in csv.DictReader(handle)}
            if any(tuple(r[c] for c in columns) in seen for r in parts[name]):
                raise ValueError(f"{file} already holds answers from {who['id']}; nothing was imported")
    if parts["A"]:
        import_sheets.append(args.output / "clinvar_annotations.csv", parts["A"])
    if parts["B"]:
        append_csv(args.output / "extraction_labels.csv", parts["B"], EXTRACTION_COLUMNS, ("sample_id", "reviewer_id"))
    if parts["C"]:
        columns = review.ANNOTATION_COLUMNS + review.OPTIONAL_ANNOTATION_COLUMNS
        append_csv(args.output / "singlecell_annotations.csv", parts["C"], columns, ("case_id", "reviewer_id"))
        review.load_annotations(args.output / "singlecell_annotations.csv")  # the protocol's strict check
    print(f"{who['id']}: part A {len(parts['A']) // 3} case(s), part B {len(parts['B'])} claim(s), "
          f"part C {len(parts['C'])} cluster(s) imported into {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
