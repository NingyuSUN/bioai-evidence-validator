"""Convert one reviewer's completed ClinVar review workbook (or labels.csv) into protocol annotations.

    uv run --with openpyxl python evaluation/clinvar_review/import_sheets.py \
        --key artifacts/clinvar-review/maintainer/key.json --reviewer-id R1 \
        --qualification "Clinical molecular geneticist, 8 years of variant curation" \
        --annotated-at 2026-10-15 --output annotations.csv reviewer_R1.xlsx

Each case becomes one annotation row per use (the statement label is shared across uses).
Rows left completely blank are skipped, so partial returns (e.g. the calibration set) can be
imported; a partly filled row is an error. Existing rows in --output are kept; the same
reviewer cannot be imported twice for a case.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import sys
from pathlib import Path

from bioevidence_validator.review import ANNOTATION_COLUMNS, OPTIONAL_ANNOTATION_COLUMNS, load_annotations, write_csv

USES = ["research_summary", "clinical_reference", "expert_reference"]
REQUIRED = ["statement_label", *USES]
CHOICES = {"statement_label": {"correct", "incorrect", "uncertain"},
           **{use: {"admitted", "review_required", "rejected"} for use in USES},
           "later_information_seen": {"no", "yes"}}
PACKET_ONLY = "review packet: ClinVar 2023-09 submissions"


def read_labels(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        sheet = load_workbook(path, read_only=True, data_only=True)["Labels"]
        rows = list(sheet.iter_rows(values_only=True))
        header = [str(h) for h in rows[0]]
        return [{h: "" if v is None else str(v).strip() for h, v in zip(header, row, strict=False)} for row in rows[1:]]
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return [{k: (v or "").strip() for k, v in row.items()} for row in csv.DictReader(handle)]


def normalize(value: str) -> str:
    return re.sub(r"[\s-]+", "_", value.strip().lower())


def convert(labels: list[dict[str, str]], key: dict, *, reviewer_id: str, qualification: str,
            annotated_at: str) -> tuple[list[dict[str, str]], int]:
    rows, errors, skipped = [], [], 0
    for number, row in enumerate(labels, start=2):
        code = row.get("case_code", "")
        filled = [c for c in [*REQUIRED, "later_information_seen"] if row.get(c)]
        if not filled:
            skipped += 1
            continue
        where = f"row {number} ({code})"
        if code not in key["cases"]:
            errors.append(f"{where}: unknown case code")
            continue
        values = {c: normalize(row.get(c, "")) for c in CHOICES}
        values["later_information_seen"] = values["later_information_seen"] or "no"
        bad = [c for c in CHOICES if values[c] not in CHOICES[c]]
        if bad:
            errors.append(f"{where}: missing or invalid {bad}")
            continue
        case = key["cases"][code]
        sources = [s.strip() for s in re.split(r"[;\n]", row.get("sources_consulted", "")) if s.strip()]
        for use in USES:
            rows.append({"annotation_id": f"{reviewer_id}:{code}:{use}", "case_id": case["case_id"],
                         "record_sha256": case["record_sha256"], "profile_id": key["profile_id"],
                         "profile_sha256": key["profile_sha256"], "requested_use": use,
                         "group_id": case["variation_id"], "split": case["split"], "reviewer_id": reviewer_id,
                         "reviewer_qualification": qualification, "annotated_at": annotated_at,
                         "mapping_label": values["statement_label"],
                         "mapping_rationale": row.get("statement_rationale", ""),
                         "admission_label": values[use], "admission_rationale": row.get(f"{use}_rationale", ""),
                         "evidence_refs_json": json.dumps(sources or [PACKET_ONLY], ensure_ascii=False),
                         "later_information_seen": values["later_information_seen"],
                         "minutes_spent": row.get("minutes_spent", "")})
    if errors:
        raise ValueError("Cannot import:\n" + "\n".join(errors[:20]))
    return rows, skipped


def append(output: Path, new: list[dict[str, str]]) -> int:
    """Add rows to an annotations file; the combined file must pass the protocol checks before it replaces the old one."""
    existing = load_annotations(output) if output.exists() else []
    columns = ANNOTATION_COLUMNS + OPTIONAL_ANNOTATION_COLUMNS
    combined = [{c: row.get(c, "") for c in columns} for row in existing] + new
    temp = output.with_suffix(".tmp.csv")
    write_csv(temp, combined, columns)
    try:
        load_annotations(temp)
    except ValueError:
        temp.unlink()
        raise
    temp.replace(output)
    return len(combined)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("labels", type=Path, help="Completed review_workbook.xlsx or labels.csv")
    parser.add_argument("--key", type=Path, required=True)
    parser.add_argument("--reviewer-id", required=True)
    parser.add_argument("--qualification", required=True)
    parser.add_argument("--annotated-at", required=True, help="ISO date or time the reviewer finished")
    parser.add_argument("--output", type=Path, required=True, help="annotations.csv to create or extend")
    args = parser.parse_args()
    dt.datetime.fromisoformat(args.annotated_at)
    key = json.loads(args.key.read_text(encoding="utf-8"))
    new, skipped = convert(read_labels(args.labels), key, reviewer_id=args.reviewer_id,
                           qualification=args.qualification, annotated_at=args.annotated_at)
    total = append(args.output, new)
    print(f"Imported {len(new) // len(USES)} case(s) from {args.reviewer_id} ({skipped} blank row(s) skipped); "
          f"{args.output} now has {total} annotation rows")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except (OSError, ValueError, KeyError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        sys.exit(3)
