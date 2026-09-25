"""Prepare blinded expert-review packets for the ClinVar case.

    uv run --with openpyxl python evaluation/clinvar_review/prepare_packets.py --output artifacts/clinvar-review

Cases: every sampled variant where the validator and NCBI's 2023-09 review status disagree on at
least one use, plus agreeing controls matched by review-status stratum. Reviewers see only the
2023-09 per-submission evidence under an opaque case code: no VariationID, SCV accession, star
rating, validator output or later ClinVar data. The maintainer key maps codes back and holds the
predictions that are scored after review; it must not be shared with reviewers.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import secrets
import shutil
from collections import Counter
from pathlib import Path

from bioevidence_validator import __version__
from bioevidence_validator.engine import RecordValidator

REPO = Path(__file__).resolve().parents[2]
CASE = REPO / "examples" / "clinvar_germline"
HERE = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location("clinvar_review_pipeline", CASE / "pipeline.py")
pipeline = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(pipeline)

RUBRIC_VERSION = "clinvar-review-rubric-1"
SELECTION_SALT = "bioai-clinvar-review-v1:"
USES = pipeline.USES
CRITERIA_TEXT = {
    "criteria provided, single submitter": "Assertion criteria provided",
    "reviewed by expert panel": "Expert panel review",
    "practice guideline": "Practice guideline",
    "no assertion criteria provided": "No assertion criteria",
    "no assertion provided": "No classification provided",
}
LABEL_COLUMNS = ["case_code", "set", "gene", "submissions", "summary",
                 "statement_label", "statement_rationale",
                 "research_summary", "research_summary_rationale",
                 "clinical_reference", "clinical_reference_rationale",
                 "expert_reference", "expert_reference_rationale",
                 "sources_consulted", "later_information_seen", "minutes_spent"]
EVIDENCE_COLUMNS = ["case_code", "submission", "submitter", "assertion_criteria", "classification",
                    "collection_method", "date_last_evaluated"]
CHOICES = {"statement_label": ["correct", "incorrect", "uncertain"],
           **{use: ["admitted", "review_required", "rejected"] for use in USES},
           "later_information_seen": ["no", "yes"]}


def evidence_rows(code: str, case: dict) -> list[dict[str, str]]:
    rows = []
    for index, sub in enumerate(sorted(case["submissions"], key=lambda s: (s["Submitter"], s["SCV"])), start=1):
        rows.append({"case_code": code, "submission": str(index), "submitter": sub["Submitter"],
                     "assertion_criteria": CRITERIA_TEXT[sub["ReviewStatus"]],
                     "classification": sub["ClinicalSignificance"], "collection_method": sub["CollectionMethod"],
                     "date_last_evaluated": sub["DateLastEvaluated"]})
    return rows


def summary(rows: list[dict[str, str]]) -> str:
    return "\n".join(f"{r['submission']}) {r['classification']} · {r['assertion_criteria']} · "
                     f"{r['collection_method']} · evaluated {r['date_last_evaluated']} · {r['submitter']}" for r in rows)


def select(controls_per_divergent: float) -> tuple[list[dict], str, str]:
    sample = pipeline.ClinVarSample()
    validator = RecordValidator(profile=CASE / "profile.yaml")
    evaluated = []
    for case in sample.cases:
        report = validator.validate(sample.record(case))
        status = {d["use"]: d["admission_status"] for d in report["use_decisions"]}
        ncbi = {use: pipeline.ncbi_expected(case["stratum"], use) for use in USES}
        divergent = any((status[u] == "admitted") != (ncbi[u] == "admitted") for u in USES)
        evaluated.append({"case": case, "record_sha256": report["input_sha256"], "validator": status,
                          "ncbi_review_status": ncbi, "role": "divergent" if divergent else "control"})
    divergent = [e for e in evaluated if e["role"] == "divergent"]
    wanted = Counter(e["case"]["stratum"] for e in divergent)
    order = lambda e: hashlib.sha256((SELECTION_SALT + e["case"]["variation_id"]).encode()).hexdigest()  # noqa: E731
    controls = []
    for stratum, count in sorted(wanted.items()):
        pool = sorted((e for e in evaluated if e["role"] == "control" and e["case"]["stratum"] == stratum), key=order)
        controls += pool[:round(count * controls_per_divergent)]
    return divergent + controls, validator.profile_sha256, sample.manifest["projection_sha256"]


def write_workbook(path: Path, labels: list[dict], evidence: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    start = book.active
    start.title = "Start"
    notes = [
        ("ClinVar evidence review", True),
        (f"Rubric: {RUBRIC_VERSION} (read RUBRIC.md before you start)", False),
        ("1. Work alone. Do not discuss cases with other reviewers until everyone has submitted.", False),
        ("2. Label the 'calibration' set first; the group discusses it before the 'test' set.", False),
        ("3. Judge only the submissions shown (ClinVar, September 2023). Do not look up the variant's "
         "current ClinVar record; if you saw later information anyway, answer 'yes' in later_information_seen.", False),
        ("4. Fill every dropdown column on the Labels sheet; add a short rationale for each label.", False),
        ("Evidence per case is on the Evidence sheet and summarized in the Labels sheet.", False),
    ]
    for row, (text, bold) in enumerate(notes, start=1):
        start.cell(row=row, column=1, value=text).font = Font(bold=bold, size=14 if bold else 11)
    start.column_dimensions["A"].width = 120

    sheet = book.create_sheet("Labels")
    sheet.append(LABEL_COLUMNS)
    for row in labels:
        sheet.append([row.get(c, "") for c in LABEL_COLUMNS])
    fill = PatternFill("solid", fgColor="FFF4CC")
    for index, column in enumerate(LABEL_COLUMNS, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.cell(row=1, column=index).font = Font(bold=True)
        sheet.column_dimensions[letter].width = {"summary": 90, "case_code": 11, "set": 12, "gene": 10,
                                                 "submissions": 11}.get(column, 22 if column in CHOICES else 36)
        if column in CHOICES or column not in LABEL_COLUMNS[:5]:
            for cell in sheet[letter][1:]:
                cell.fill = fill
        if column in CHOICES:
            rule = DataValidation(type="list", formula1='"' + ",".join(CHOICES[column]) + '"', allow_blank=True)
            rule.add(f"{letter}2:{letter}{len(labels) + 1}")
            sheet.add_data_validation(rule)
    for cell in sheet["E"][1:]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "F2"

    ev = book.create_sheet("Evidence")
    ev.append(EVIDENCE_COLUMNS)
    for row in evidence:
        ev.append([row[c] for c in EVIDENCE_COLUMNS])
    for index in range(1, len(EVIDENCE_COLUMNS) + 1):
        ev.cell(row=1, column=index).font = Font(bold=True)
        ev.column_dimensions[ev.cell(row=1, column=index).column_letter].width = 30
    ev.freeze_panes = "B2"
    book.save(path)


def write_csv(path: Path, rows: list[dict], columns: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def prepare(output: Path, *, salt: str, controls_per_divergent: float = 1.0, calibration: int = 20) -> dict:
    if output.exists():
        raise ValueError("Use a new output directory")
    selected, profile_sha256, source_sha256 = select(controls_per_divergent)
    coded = {}
    for entry in selected:
        code = "C" + hashlib.sha256((salt + entry["case"]["variation_id"]).encode()).hexdigest()[:6].upper()
        if code in coded:
            raise ValueError("Case-code collision; choose another salt")
        coded[code] = entry
    ordered = sorted(coded)
    labels, evidence, key_cases, predictions = [], [], {}, []
    for position, code in enumerate(ordered):
        entry, case = coded[code], coded[code]["case"]
        split = "development" if position < calibration else "test"
        rows = evidence_rows(code, case)
        evidence += rows
        labels.append({"case_code": code, "set": "calibration" if split == "development" else "test",
                       "gene": case["gene"], "submissions": str(len(rows)), "summary": summary(rows)})
        case_id = "clinvar:" + case["variation_id"]
        key_cases[code] = {"case_id": case_id, "variation_id": case["variation_id"], "stratum": case["stratum"],
                           "role": entry["role"], "split": split, "record_sha256": entry["record_sha256"],
                           "validator": entry["validator"], "ncbi_review_status": entry["ncbi_review_status"]}
        for use in USES:
            for method, predicted in [("validator", entry["validator"][use]),
                                      ("ncbi_review_status", entry["ncbi_review_status"][use])]:
                predictions.append({"case_id": case_id, "requested_use": use, "record_sha256": entry["record_sha256"],
                                    "profile_sha256": profile_sha256, "method": method,
                                    "predicted_status": predicted, "subset": entry["role"]})

    packet, private = output / "reviewer_packet", output / "maintainer"
    packet.mkdir(parents=True)
    private.mkdir()
    write_workbook(packet / "review_workbook.xlsx", labels, evidence)
    write_csv(packet / "labels.csv", labels, LABEL_COLUMNS)
    write_csv(packet / "evidence.csv", evidence, EVIDENCE_COLUMNS)
    shutil.copyfile(HERE / "RUBRIC.md", packet / "RUBRIC.md")

    counts = {"cases": len(ordered), "divergent": sum(e["role"] == "divergent" for e in coded.values()),
              "controls": sum(e["role"] == "control" for e in coded.values()),
              "calibration": min(calibration, len(ordered)), "test": max(0, len(ordered) - calibration),
              "strata": dict(sorted(Counter(e["case"]["stratum"] for e in coded.values()).items()))}
    key = {"warning": "Maintainer only. Do not share with reviewers before all reviews are submitted.",
           "rubric_version": RUBRIC_VERSION, "code_salt": salt, "selection_salt": SELECTION_SALT,
           "validator_version": __version__, "profile_id": "clinvar-germline", "profile_sha256": profile_sha256,
           "source_projection_sha256": source_sha256, "counts": counts, "cases": key_cases}
    (private / "key.json").write_text(json.dumps(key, indent=2) + "\n", encoding="utf-8", newline="\n")
    write_csv(private / "predictions.csv", predictions,
              ["case_id", "requested_use", "record_sha256", "profile_sha256", "method", "predicted_status", "subset"])
    template = json.loads((REPO / "evaluation/gold_standard/manifest.template.json").read_text(encoding="utf-8"))
    manifest = {**template, "dataset_id": "clinvar-germline-expert-review", "status": "draft",
                "protocol_version": RUBRIC_VERSION,
                "source_snapshots": [{"source": "ClinVar 2023-09 germline sample (examples/clinvar_germline)",
                                      "sha256": source_sha256}],
                "profiles": [{"id": "clinvar-germline", "sha256": profile_sha256}], "selection": counts}
    (private / "manifest.draft.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8", newline="\n")
    return counts


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--salt", help="Secret used to derive case codes (random if omitted; kept in the key)")
    parser.add_argument("--controls-per-divergent", type=float, default=1.0)
    parser.add_argument("--calibration", type=int, default=20, help="Cases in the calibration (development) set")
    args = parser.parse_args()
    result = prepare(args.output, salt=args.salt or secrets.token_hex(16),
                     controls_per_divergent=args.controls_per_divergent, calibration=args.calibration)
    print(json.dumps(result, indent=2))
